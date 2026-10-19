"""Import guarantees, checked on the real workbooks in data/imports:  python -m unittest pipeline.test_import -v
 - importing twice changes nothing
 - a re-upload with Actual typed in settles the Log but never rewrites its frozen projections
 - a workbook for an older slate updates only that slate's Log, never the live inputs
"""
import json, os, shutil, tempfile, unittest
import openpyxl
from . import import_xlsx as I

SLATE = "2026-10-16"
FPS, MOBA = f"data/imports/{SLATE}/fps.xlsx", f"data/imports/{SLATE}/moba.xlsx"


def run(site, imports, fps=FPS, moba=MOBA, slate=None):
    import sys
    argv = sys.argv
    sys.argv = ["import", "--fps", fps] + (["--moba", moba] if moba else []) + ["--site", site, "--imports", imports] + (["--slate", slate] if slate else [])
    try:
        I.main()
    finally:
        sys.argv = argv


def snap(site):
    out = {}
    for root, _, files in os.walk(site):
        for f in files:
            p = os.path.join(root, f)
            out[os.path.relpath(p, site)] = open(p, "rb").read()
    return out


def edited(src, dst, fn):
    wb = openpyxl.load_workbook(src, data_only=True)
    fn(wb)
    wb.save(dst)


class ImportTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.site, self.imports = f"{self.tmp}/site", f"{self.tmp}/imports"

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def test_idempotent(self):
        run(self.site, self.imports)
        a = snap(self.site)
        run(self.site, self.imports)
        self.assertEqual(a, snap(self.site))

    def test_settlement_adds_actuals_and_keeps_frozen_values(self):
        run(self.site, self.imports)
        log0 = json.load(open(f"{self.site}/data/log/{SLATE}.json"))

        def settle(wb):
            ws = wb["Log"]
            ws["O2"] = 33                                              # Actual
            ws["P2"] = "OVER"
            ws["H2"] = 999.0                                           # a rebuilt workbook must not be able to rewrite the frozen Proj

        fps2 = f"{self.tmp}/fps2.xlsx"
        edited(FPS, fps2, settle)
        run(self.site, self.imports, fps=fps2)
        log1 = json.load(open(f"{self.site}/data/log/{SLATE}.json"))
        self.assertEqual(log1["fps"][0]["Actual"], 33)
        self.assertEqual(log1["fps"][0]["Proj"], log0["fps"][0]["Proj"])
        self.assertEqual(sum(1 for r in log1["fps"] if r.get("Actual") is not None), 1)
        self.assertEqual(len(log1["fps"]), len(log0["fps"]))
        # a later upload that has no Actual must not erase the settled one
        run(self.site, self.imports)
        self.assertEqual(json.load(open(f"{self.site}/data/log/{SLATE}.json"))["fps"][0]["Actual"], 33)

    def test_older_slate_only_touches_its_log(self):
        run(self.site, self.imports)
        before = snap(self.site)

        def older(wb):
            ws = wb["Config"]
            for r in range(1, 30):
                if ws.cell(r, 1).value == "SLATE_DATE":
                    ws.cell(r, 2).value = "2026-10-15"

        f2, m2 = f"{self.tmp}/f.xlsx", f"{self.tmp}/m.xlsx"
        edited(FPS, f2, older)
        edited(MOBA, m2, older)
        run(self.site, self.imports, fps=f2, moba=m2)
        after = snap(self.site)
        for name in ("data/fps.json", "data/moba.json", "test/expected.json"):
            self.assertEqual(before[name], after[name], name)
        meta = json.loads(after["data/meta.json"])
        self.assertEqual(meta["slate"], SLATE)
        self.assertEqual(meta["slates"], ["2026-10-15", SLATE])
        self.assertIn("data/log/2026-10-15.json", after)


class Refusals(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.site, self.imports = f"{self.tmp}/site", f"{self.tmp}/imports"

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def relabel(self, src, dst, slate):
        def f(wb):
            ws = wb["Config"]
            for r in range(1, 40):
                if ws.cell(r, 1).value == "SLATE_DATE":
                    ws.cell(r, 2).value = slate
        edited(src, dst, f)

    def test_workbooks_for_different_slates_import_nothing(self):
        m2 = f"{self.tmp}/m.xlsx"
        self.relabel(MOBA, m2, "2026-10-17")
        with self.assertRaises(SystemExit) as e:
            run(self.site, self.imports, moba=m2)
        self.assertIn("different slates", str(e.exception))
        self.assertFalse(os.path.exists(self.site))

    def test_yesterdays_workbook_in_todays_folder_is_refused(self):
        with self.assertRaises(SystemExit) as e:
            run(self.site, self.imports, slate="2026-10-17")
        self.assertIn("2026-10-17", str(e.exception))
        self.assertIn(SLATE, str(e.exception))
        self.assertFalse(os.path.exists(self.site))

    def test_a_workbook_saved_without_results_is_refused(self):
        stripped = f"{self.tmp}/fps_stripped.xlsx"
        openpyxl.load_workbook(FPS).save(stripped)                      # formulas kept, stored results dropped (what a script-built workbook looks like)
        with self.assertRaises(SystemExit) as e:
            run(self.site, self.imports, fps=stripped)
        self.assertIn("no stored results", str(e.exception))
        self.assertFalse(os.path.exists(self.site))

    def test_the_matching_folder_is_accepted(self):
        run(self.site, self.imports, slate=SLATE)
        self.assertTrue(os.path.exists(f"{self.site}/data/meta.json"))


class CS2OnlyDay(unittest.TestCase):
    """A day with no LoL/Dota workbook builds with no MOBA legs, and the board reads back without one."""
    def test_fps_only_import(self):
        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp, ignore_errors=True)
        site, imports = f"{tmp}/site", f"{tmp}/imports"
        run(site, imports, moba=None, slate=SLATE)
        meta = json.load(open(f"{site}/data/meta.json"))
        self.assertEqual((meta["slate"], meta["moba_legs"], meta["moba_matches"], meta["fps_legs"] > 0), (SLATE, 0, 0, True))
        self.assertEqual(json.load(open(f"{site}/data/moba.json"))["legs"], [])
        self.assertFalse(os.path.exists(f"{imports}/{SLATE}/moba.xlsx"))
        plays = f"{tmp}/plays/{SLATE}.json"                                     # the play-rule selections are frozen at import, beside the imports folder
        self.assertTrue(os.path.exists(plays))
        frozen = json.load(open(plays))
        self.assertEqual(frozen["rules"]["all_under"]["n"], fps_legs := json.load(open(f"{site}/data/meta.json"))["fps_legs"])
        self.assertTrue(all(l["side"] == "UNDER" for l in frozen["rules"]["all_under"]["legs"]))
        self.assertTrue(all(l["side"] in ("OVER", "UNDER") for l in frozen["rules"]["agree10"]["legs"]))
        mtime = os.path.getmtime(plays)
        run(site, imports, moba=None, slate=SLATE)                              # importing again never rewrites a frozen file
        self.assertEqual(os.path.getmtime(plays), mtime)
        self.assertGreater(fps_legs, 0)
        self.assertEqual(json.load(open(f"{site}/data/log/{SLATE}.json"))["moba"], [])
        os.makedirs(f"{tmp}/s2/data")
        json.dump(dict(slate="2099-01-01"), open(f"{tmp}/s2/data/meta.json", "w"))               # a different slate is live: this one is read back from its workbooks
        from .results import load_board
        _, moba, where = load_board(SLATE, f"{tmp}/s2", imports)
        self.assertEqual((where, moba["legs"]), ("workbooks", []))


if __name__ == "__main__":
    unittest.main()
